import sys
import os, ssl
import urllib.request
from bs4 import BeautifulSoup, SoupStrainer
import re
import requests
from pprint import pprint
from selenium import webdriver
import xlsxwriter
from openpyxl import load_workbook
import time
import urllib.request
from html_table_parser.parser import HTMLTableParser
import pandas as pd
from datetime import date

ssl._create_default_https_context = ssl._create_unverified_context
response = urllib.request.urlopen('https://www.python.org')
print(response.read().decode('utf-8'))

def url_get_contents(url):
    # Opens a website and read its
    # binary contents (HTTP Response Body)

    #making request to the website
    req = urllib.request.Request(url=url)
    f = urllib.request.urlopen(req)

    #reading contents of the website
    return f.read()

Path=r"F:\Projects\NFL Stats\ "

writer = pd.ExcelWriter(Path + "NFL Stats.xlsx", engine = 'xlsxwriter')


#workbook = xlsxwriter.Workbook(Path + 'NFL Stats Test.xlsx')


#This section creates a dataframe of the offensive PASSING stats for all NFL Teams
NFL_Passing_url = "https://www.nfl.com/stats/team-stats/offense/passing/2021/reg/all"
xhtml = url_get_contents(NFL_Passing_url).decode('utf-8')
p = HTMLTableParser()
p.feed(xhtml)
df1 = pd.DataFrame(p.tables[0])
df1.columns = ['Team','Att','Cmp','Cmp %','Yds/Att','Pass Yds','TD','INT','Rate','1st','1st%','20','40','Lng','Sck','SckY']
df1 = df1.drop([0])
df1["Att"] = pd.to_numeric(df1["Att"])
df1["Cmp"] = pd.to_numeric(df1["Cmp"])
df1["Cmp %"] = pd.to_numeric(df1["Cmp %"])
df1["Yds/Att"] = pd.to_numeric(df1["Yds/Att"])
df1["Pass Yds"] = pd.to_numeric(df1["Pass Yds"])
df1["TD"] = pd.to_numeric(df1["TD"])
df1["INT"] = pd.to_numeric(df1["INT"])
df1["Rate"] = pd.to_numeric(df1["Rate"])
df1["1st"] = pd.to_numeric(df1["1st"])
df1["1st%"] = pd.to_numeric(df1["1st%"])
df1["20"] = pd.to_numeric(df1["20"])
df1["40"] = pd.to_numeric(df1["40"])
df1["Lng"] = [s.replace("T","") for s in df1["Lng"]]
df1["Lng"] = pd.to_numeric(df1["Lng"])
df1["Sck"] = pd.to_numeric(df1["Sck"])
df1["SckY"] = pd.to_numeric(df1["SckY"])
df1['Att Rank'] = round(df1["Att"].rank(ascending=False, method='first'))
df1['Cmp % Rank'] = df1["Cmp %"].rank(ascending=False, method='first')
df1["Pass Yds Rank"] =df1["Pass Yds"].rank(ascending=False, method='first')
df1["TD Rank"] =df1["TD"].rank(ascending=False, method='first')
df1["INT Rank"] =df1["INT"].rank(ascending=False, method='first')
df1["Rate Rank"] =df1["Rate"].rank(ascending=False, method='first')
df1 = df1[['Team','Att','Att Rank','Cmp','Cmp %','Cmp % Rank','Yds/Att','Pass Yds',"Pass Yds Rank",'TD',"TD Rank",'INT',"INT Rank",'Rate',"Rate Rank",'1st','1st%','20','40','Lng','Sck','SckY']]


df1.loc[-1] = ['Team','Att','Att Rank','Cmp','Cmp %','Cmp % Rank','Yds/Att','Pass Yds',"Pass Yds Rank",'TD',"TD Rank",'INT',"INT Rank",'Rate',"Rate Rank",'1st','1st%','20','40','Lng','Sck','SckY']  # adding a row
df1.index = df1.index + 1  # shifting index
df1.sort_index(inplace=True)

#ws1 = workbook.add_worksheet('Offense Passing')
df1.to_excel(writer,sheet_name="Offense Passing" ,index=False, header=False)


NFL_Rushing_url = "https://www.nfl.com/stats/team-stats/offense/rushing/2021/reg/all"
xhtml = url_get_contents(NFL_Rushing_url).decode('utf-8')
p = HTMLTableParser()
p.feed(xhtml)
df2 = pd.DataFrame(p.tables[0])
df2.columns = ['Team','Att','Rush Yds','YPC','TD','20','40','Lng','Rush 1st','Rush 1st %','Rush Fumble']
df2 = df2.drop([0])
df2["Att"] = pd.to_numeric(df2["Att"])
df2["Rush Yds"] = pd.to_numeric(df2["Rush Yds"])
df2["YPC"] = pd.to_numeric(df2["YPC"])
df2["TD"] = pd.to_numeric(df2["TD"])
df2["20"] = pd.to_numeric(df2["20"])
df2["40"] = pd.to_numeric(df2["40"])
df2["Rush 1st"] = pd.to_numeric(df2["Rush 1st"])
df2["Rush 1st %"] = pd.to_numeric(df2["Rush 1st %"])
df2["Rush Fumble"] = pd.to_numeric(df2["Rush Fumble"])
df2["Lng"] = [s.replace("T","") for s in df2["Lng"]]
df2["Lng"] = pd.to_numeric(df2["Lng"])
df2['Att Rank'] = round(df2["Att"].rank(ascending=False, method='first'))
df2['Rush Yds Rank'] = df2["Rush Yds"].rank(ascending=False, method='first')
df2["TD Rank"] =df2["TD"].rank(ascending=False, method='first')
df2 = df2[['Team','Att','Att Rank','Rush Yds','Rush Yds Rank','YPC','TD',"TD Rank",'20','40','Lng','Rush 1st','Rush 1st %','Rush Fumble']]


df2.loc[-1] = ['Team','Att','Att Rank','Rush Yds','Rush Yds Rank','YPC','TD',"TD Rank",'20','40','Lng','Rush 1st','Rush 1st %','Rush Fumble']  # adding a row
df2.index = df2.index + 1  # shifting index
df2.sort_index(inplace=True)
#ws2 = writer.add_worksheet("Offense Rushing")
df2.to_excel(writer,sheet_name="Offense Rushing",index=False, header=False)
#filename = Path + "Offense_Rushing.xlsx"

NFL_PassingD_url = "https://www.nfl.com/stats/team-stats/defense/passing/2021/reg/all"
xhtml = url_get_contents(NFL_PassingD_url).decode('utf-8')
p = HTMLTableParser()
p.feed(xhtml)
df3 = pd.DataFrame(p.tables[0])
df3.columns = ['Team','Att','Cmp','Cmp %','Yds/Att','Pass Yds','TD','INT','Rate','1st','1st%','20','40','Lng','Sck']
df3 = df3.drop([0])
df3["Att"] = pd.to_numeric(df3["Att"])
df3["Cmp"] = pd.to_numeric(df3["Cmp"])
df3["Cmp %"] = pd.to_numeric(df3["Cmp %"])
df3["Yds/Att"] = pd.to_numeric(df3["Yds/Att"])
df3["Pass Yds"] = pd.to_numeric(df3["Pass Yds"])
df3["TD"] = pd.to_numeric(df3["TD"])
df3["INT"] = pd.to_numeric(df3["INT"])
df3["Rate"] = pd.to_numeric(df3["Rate"])
df3["1st"] = pd.to_numeric(df3["1st"])
df3["1st%"] = pd.to_numeric(df3["1st%"])
df3["20"] = pd.to_numeric(df3["20"])
df3["40"] = pd.to_numeric(df3["40"])
df3["Lng"] = [s.replace("T","") for s in df3["Lng"]]
df3["Lng"] = pd.to_numeric(df3["Lng"])
df3["Sck"] = pd.to_numeric(df3["Sck"])
df3['Att Rank'] = round(df3["Att"].rank(ascending=True, method='first'))
df3['Cmp % Rank'] = df3["Cmp %"].rank(ascending=True, method='first')
df3["Pass Yds Rank"] =df3["Pass Yds"].rank(ascending=True, method='first')
df3["TD Rank"] =df3["TD"].rank(ascending=True, method='first')
df3["INT Rank"] =df3["INT"].rank(ascending=True, method='first')
df3["Rate Rank"] =df3["Rate"].rank(ascending=True, method='first')
df3 = df3[['Team','Att','Att Rank','Cmp','Cmp %','Cmp % Rank','Yds/Att','Pass Yds',"Pass Yds Rank",'TD',"TD Rank",'INT',"INT Rank",'Rate',"Rate Rank",'1st','1st%','20','40','Lng','Sck']]


df3.loc[-1] = ['Team','Att','Att Rank','Cmp','Cmp %','Cmp % Rank','Yds/Att','Pass Yds',"Pass Yds Rank",'TD',"TD Rank",'INT',"INT Rank",'Rate',"Rate Rank",'1st','1st%','20','40','Lng','Sck']  # adding a row
df3.index = df3.index + 1  # shifting index
df3.sort_index(inplace=True)
df3.to_excel(writer,sheet_name="Defense Passing",index=False, header=False)
#filename = Path + "Defense_Passing.xlsx"

NFL_RushingD_url = "https://www.nfl.com/stats/team-stats/defense/rushing/2021/reg/all"
xhtml = url_get_contents(NFL_RushingD_url).decode('utf-8')
p = HTMLTableParser()
p.feed(xhtml)
df4 = pd.DataFrame(p.tables[0])
df4.columns = ['Team','Att','Rush Yds','YPC','TD','20','40','Lng','Rush 1st','Rush 1st %','Rush Fumble']
df4 = df4.drop([0])
df4["Att"] = pd.to_numeric(df4["Att"])
df4["Rush Yds"] = pd.to_numeric(df4["Rush Yds"])
df4["YPC"] = pd.to_numeric(df4["YPC"])
df4["TD"] = pd.to_numeric(df4["TD"])
df4["20"] = pd.to_numeric(df4["20"])
df4["40"] = pd.to_numeric(df4["40"])
df4["Rush 1st"] = pd.to_numeric(df4["Rush 1st"])
df4["Rush 1st %"] = pd.to_numeric(df4["Rush 1st %"])
df4["Rush Fumble"] = pd.to_numeric(df4["Rush Fumble"])
df4["Lng"] = [s.replace("T","") for s in df4["Lng"]]
df4["Lng"] = pd.to_numeric(df4["Lng"])
df4['Att Rank'] = round(df4["Att"].rank(ascending=True, method='first'))
df4['Rush Yds Rank'] = df4["Rush Yds"].rank(ascending=True, method='first')
df4["TD Rank"] =df4["TD"].rank(ascending=True, method='first')
df4 = df4[['Team','Att','Att Rank','Rush Yds','Rush Yds Rank','YPC','TD',"TD Rank",'20','40','Lng','Rush 1st','Rush 1st %','Rush Fumble']]


df4.loc[-1] = ['Team','Att','Att Rank','Rush Yds','Rush Yds Rank','YPC','TD',"TD Rank",'20','40','Lng','Rush 1st','Rush 1st %','Rush Fumble']  # adding a row
df4.index = df4.index + 1  # shifting index
df4.sort_index(inplace=True)
df4.to_excel(writer,sheet_name="Defense Rushing",index=False, header=False)
#filename = Path + "Defense_Rushing.xlsx"


espn_Offense_url = "https://www.espn.com/nfl/stats/team"
xhtml = url_get_contents(espn_Offense_url).decode('utf-8')
p = HTMLTableParser()
p.feed(xhtml)
df5 = pd.DataFrame(p.tables[0])
df6 = pd.DataFrame(p.tables[1])
df5 = df5.drop([0])
df5 = df5.drop([1])
df5.loc[-1] = ['Team']
df5.index = df5.index +1
df5.sort_index(inplace=True)
df6 = df6.drop([0])
df6 = df6.drop([1])
df6.columns = ['GP','Total Yards','Total Yds/G','Passing Total Yds','Passing Yds/G','Rushing Total Yds','Rushing Yds/G','Total Points','Pts/G']
df6["GP"] = pd.to_numeric(df6["GP"])
df6["Total Yards"] = [s.replace(",","") for s in df6["Total Yards"]]
df6["Total Yards"] = pd.to_numeric(df6["Total Yards"])
df6["Total Yds/G"] = pd.to_numeric(df6["Total Yds/G"])
df6["Passing Total Yds"] = [s.replace(",","") for s in df6["Passing Total Yds"]]
df6['Passing Total Yds'] = pd.to_numeric(df6['Passing Total Yds'])
df6['Passing Yds/G'] = pd.to_numeric(df6['Passing Yds/G'])
df6["Rushing Total Yds"] = [s.replace(",","") for s in df6["Rushing Total Yds"]]
df6['Rushing Total Yds'] = pd.to_numeric(df6['Rushing Total Yds'])
df6['Rushing Yds/G'] = pd.to_numeric(df6['Rushing Yds/G'])
df6["Total Points"] = [s.replace(",","") for s in df6["Total Points"]]
df6['Total Points'] = pd.to_numeric(df6['Total Points'])
df6['Pts/G'] = pd.to_numeric(df6['Pts/G'])
df6['Rush Yds/G Rank'] = df6["Rushing Yds/G"].rank(ascending=False, method='first')
df6["Passing Yds/G Rank"] =df6["Passing Yds/G"].rank(ascending=False, method='first')
df6["Total Yds/G Rank"] =df6["Total Yds/G"].rank(ascending=False, method='first')
df6["Pts/G Rank"] =df6["Pts/G"].rank(ascending=False, method='first')
df6 = df6[['GP','Total Yards','Total Yds/G',"Total Yds/G Rank",'Passing Total Yds','Passing Yds/G',"Passing Yds/G Rank",
    'Rushing Total Yds','Rushing Yds/G','Rush Yds/G Rank','Total Points','Pts/G',"Pts/G Rank"]]

df6.loc[-1] = ['GP','Total Yards','Total Yds/G',"Total Yds/G Rank",'Passing Total Yds','Passing Yds/G',"Passing Yds/G Rank",
    'Rushing Total Yds','Rushing Yds/G','Rush Yds/G Rank','Total Points','Pts/G',"Pts/G Rank"]
df6.index = df6.index +1
df6.sort_index(inplace=True)
frames1 = [df5, df6]
results1 = pd.concat(frames1, axis=1)
results1.to_excel(writer,sheet_name="Offense ESPN",index=False, header=False)
#filename = Path + "Offense_ESPN.xlsx"


espn_Defense_url = "https://www.espn.com/nfl/stats/team/_/view/defense"
xhtml = url_get_contents(espn_Defense_url).decode('utf-8')
p = HTMLTableParser()
p.feed(xhtml)
df7 = pd.DataFrame(p.tables[0])
df7 = df7.drop([0])
df7 = df7.drop([1])
df7.loc[-1] = ['Team']
df7.index = df7.index +1
df7.sort_index(inplace=True)
df8 = pd.DataFrame(p.tables[1])
df8 = df8.drop([0])
df8 = df8.drop([1])
df8.columns = ['GP','Total Yards','Total Yds/G','Passing Total Yds','Passing Yds/G','Rushing Total Yds','Rushing Yds/G','Total Points','Pts/G']
df8["GP"] = pd.to_numeric(df8["GP"])
df8["Total Yards"] = [s.replace(",","") for s in df8["Total Yards"]]
df8["Total Yards"] = pd.to_numeric(df8["Total Yards"])
df8["Total Yds/G"] = pd.to_numeric(df8["Total Yds/G"])
df8["Passing Total Yds"] = [s.replace(",","") for s in df8["Passing Total Yds"]]
df8['Passing Total Yds'] = pd.to_numeric(df8['Passing Total Yds'])
df8['Passing Yds/G'] = pd.to_numeric(df8['Passing Yds/G'])
df8["Rushing Total Yds"] = [s.replace(",","") for s in df8["Rushing Total Yds"]]
df8['Rushing Total Yds'] = pd.to_numeric(df8['Rushing Total Yds'])
df8['Rushing Yds/G'] = pd.to_numeric(df8['Rushing Yds/G'])
df8["Total Points"] = [s.replace(",","") for s in df8["Total Points"]]
df8['Total Points'] = pd.to_numeric(df8['Total Points'])
df8['Pts/G'] = pd.to_numeric(df8['Pts/G'])
df8['Rush Yds/G Rank'] = df8["Rushing Yds/G"].rank(ascending=True, method='first')
df8["Passing Yds/G Rank"] =df8["Passing Yds/G"].rank(ascending=True, method='first')
df8["Total Yds/G Rank"] =df8["Total Yds/G"].rank(ascending=True, method='first')
df8["Pts/G Rank"] =df8["Pts/G"].rank(ascending=True, method='first')
df8 = df8[['GP','Total Yards','Total Yds/G',"Total Yds/G Rank",'Passing Total Yds','Passing Yds/G',"Passing Yds/G Rank",
    'Rushing Total Yds','Rushing Yds/G','Rush Yds/G Rank','Total Points','Pts/G',"Pts/G Rank"]]

df8.loc[-1] = ['GP','Total Yards','Total Yds/G',"Total Yds/G Rank",'Passing Total Yds','Passing Yds/G',"Passing Yds/G Rank",
    'Rushing Total Yds','Rushing Yds/G','Rush Yds/G Rank','Total Points','Pts/G',"Pts/G Rank"]
df8.index = df8.index +1
df8.sort_index(inplace=True)
frames2 = [df7, df8]
results2 = pd.concat(frames2, axis=1)
results2.to_excel(writer,sheet_name="Defense ESPN",index=False, header=False)
#filename = Path + "Defense_ESPN.xlsx"


#This section will apply the conditional formatting to the sheets
#filename = Path + 'NFL Stats.xlsx'

workbook = writer.book
ws1 = writer.sheets['Offense Passing']
ws2 = writer.sheets['Offense Rushing']
ws3 = writer.sheets['Defense Passing']
ws4 = writer.sheets['Defense Rushing']
ws5 = writer.sheets['Offense ESPN']
ws6 = writer.sheets['Defense ESPN']

ws1.conditional_format('C2:C33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws1.conditional_format('F2:F33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws1.conditional_format('I2:I33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws1.conditional_format('K2:K33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws1.conditional_format('M2:M33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws1.conditional_format('O2:O33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws1.autofilter('A1:A33')


ws2.conditional_format('C2:C33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws2.conditional_format('E2:E33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws2.conditional_format('H2:H33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws2.autofilter('A1:A33')


ws3.conditional_format('C2:C33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws3.conditional_format('F2:F33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws3.conditional_format('I2:I33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws3.conditional_format('K2:K33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws3.conditional_format('M2:M33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws3.conditional_format('O2:O33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws3.autofilter('A1:A33')


ws4.conditional_format('C2:C33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws4.conditional_format('E2:E33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws4.conditional_format('H2:H33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws4.autofilter('A1:A33')


ws5.conditional_format('E2:E33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws5.conditional_format('H2:H33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws5.conditional_format('K2:K33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws5.conditional_format('N2:N33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws5.autofilter('A1:A33')


ws6.conditional_format('E2:E33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws6.conditional_format('H2:H33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws6.conditional_format('K2:K33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws6.conditional_format('N2:N33', {'type': '3_color_scale', 'min_color':'#008000', 'mid_color':'#FFFF00', 'max_color':'#FF0000'})
ws6.autofilter('A1:A33')





writer.close()
